#!/usr/bin/env python3
"""Render the supported Markdown guide to Word; preserve unchanged files."""

import argparse
from io import BytesIO
from zipfile import BadZipFile, ZipFile
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
parser = argparse.ArgumentParser(description='从 Markdown 同步 AutoPPT Word 使用说明')
parser.add_argument('--check', action='store_true', help='只检查内容是否同步，不写入文件')
args = parser.parse_args()
root = Path(__file__).resolve().parents[1]
source = root / 'docs/同事安装与使用说明.md'
out = root / 'docs/AutoPPT-安装与使用说明.docx'
doc = Document()
for style in doc.styles:
    for borders in style.element.findall('.//' + qn('w:pBdr')):
        borders.getparent().remove(borders)
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.68)
sec.left_margin = sec.right_margin = Inches(0.78)
sec.header_distance = sec.footer_distance = Inches(0.3)
for grid in sec._sectPr.findall(qn('w:docGrid')):
    sec._sectPr.remove(grid)
for style in doc.styles:
    if style.type == 1:
        snap = OxmlElement('w:snapToGrid')
        snap.set(qn('w:val'), '0')
        style.element.get_or_add_pPr().append(snap)

def font(style, name='Arial', east='Arial Unicode MS', size=11):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:eastAsia'), east)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number', 'Header', 'Footer']:
    font(doc.styles[name])
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = Pt(17)
normal.space_after = Pt(5)
normal.widow_control = True
for name, size, before, after in [('Title', 24, 0, 12), ('Heading 1', 17, 14, 8), ('Heading 2', 12.5, 10, 5)]:
    s = doc.styles[name]
    font(s, size=size)
    s.font.bold = name != 'Title'
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.line_spacing = Pt(size + 6)
font(doc.styles['Header'], size=9)
font(doc.styles['Footer'], size=9)
header = sec.header.paragraphs[0]
header.text = 'AutoPPT  安装与使用说明'
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('第 ')
fld = OxmlElement('w:fldSimple')
fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)
footer.add_run(' 页')
doc.core_properties.title = 'AutoPPT 安装与使用说明'
doc.core_properties.subject = '安装、模型配置、首次测试与日常维护'
doc.core_properties.author = 'AutoPPT'
doc.core_properties.comments = ''

def inline(p, text):
    pat = '(\\*\\*[^*]+\\*\\*|`[^`]+`|\\[[^\\]]+\\]\\(https?://[^)]+\\))'
    for part in re.split(pat, text):
        if not part:
            continue
        link = re.fullmatch('\\[([^\\]]+)\\]\\(([^)]+)\\)', part)
        if link:
            h = OxmlElement('w:hyperlink')
            h.set(qn('r:id'), p.part.relate_to(link[2], RT.HYPERLINK, is_external=True))
            r = OxmlElement('w:r')
            pr = OxmlElement('w:rPr')
            col = OxmlElement('w:color')
            col.set(qn('w:val'), '245C85')
            pr.append(col)
            under = OxmlElement('w:u')
            under.set(qn('w:val'), 'single')
            pr.append(under)
            r.append(pr)
            t = OxmlElement('w:t')
            t.text = link[1]
            r.append(t)
            h.append(r)
            p._p.append(h)
        else:
            bold = part.startswith('**') and part.endswith('**')
            code = part.startswith('`') and part.endswith('`')
            r = p.add_run(part[2:-2] if bold else part[1:-1] if code else part)
            r.bold = bold
            if code:
                r.font.name = 'Menlo'
                r.font.size = Pt(10)

def table(rows):
    t = doc.add_table(rows=1, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    widths = [1.6, 3.25, 2.05]
    for c, w in zip(t.columns, widths):
        c.width = Inches(w)
    for index, row in enumerate(rows):
        cells = t.rows[0].cells if index == 0 else t.add_row().cells
        for cell, text, w in zip(cells, row, widths):
            cell.width = Inches(w)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                b = OxmlElement('w:' + edge)
                b.set(qn('w:val'), 'single')
                b.set(qn('w:sz'), '4')
                b.set(qn('w:color'), 'D9D9D9')
                borders.append(b)
            tcpr.append(borders)
            margin = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:w'), '100')
                e.set(qn('w:type'), 'dxa')
                margin.append(e)
            tcpr.append(margin)
            if index == 0:
                sh = OxmlElement('w:shd')
                sh.set(qn('w:fill'), 'EDEDED')
                tcpr.append(sh)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(16)
            inline(p, text)
            for r in p.runs:
                r.font.size = Pt(10.5)
                r.bold = index == 0
        trpr = t.rows[index]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if index == 0:
            trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
lines = source.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1
        continue
    if line.startswith('```'):
        code = []
        i += 1
        while i < len(lines) and (not lines[i].startswith('```')):
            code.append(lines[i])
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.line_spacing = Pt(16)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.keep_together = True
        r = p.add_run('\n'.join(code))
        r.font.name = 'Menlo'
        r.font.size = Pt(10.5)
        i += 1
        continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            row = [c.strip() for c in lines[i].strip('|').split('|')]
            if not all((re.fullmatch(':?-+:?', c) for c in row)):
                rows.append(row)
            i += 1
        table(rows)
        continue
    m = re.match('^(#{1,3}) (.*)', line)
    if m:
        n = len(m[1])
        style = {1: 'Title', 2: 'Heading 1', 3: 'Heading 2'}[n]
        p = doc.add_paragraph(style=style)
        inline(p, m[2])
        if n == 2 and (not m[2].startswith('1 ')):
            p.paragraph_format.page_break_before = True
        if n == 3 and m[2] == '做一次完整测试':
            p.paragraph_format.page_break_before = True
    elif line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        inline(p, line[2:])
    elif re.match('^\\d+\\. ', line):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        inline(p, line)
    else:
        p = doc.add_paragraph()
        inline(p, line)
    i += 1

def package_contents(data):
    with ZipFile(BytesIO(data)) as package:
        return {name: package.read(name) for name in package.namelist()}
buffer = BytesIO()
doc.save(buffer)
generated = buffer.getvalue()
try:
    matches = out.exists() and package_contents(out.read_bytes()) == package_contents(generated)
except (OSError, ValueError, BadZipFile):
    matches = False
if args.check:
    if not matches:
        parser.exit(1, 'Word 使用说明与 Markdown 不一致，请先运行 scripts/build-guide.py。\n')
    print('Markdown 与 Word 内容一致。')
elif matches:
    print('Word 内容未变化，保留现有文件。')
else:
    out.parent.mkdir(exist_ok=True)
    temporary = out.with_suffix('.docx.tmp')
    temporary.write_bytes(generated)
    temporary.replace(out)
    print(f'已生成 {out}')
